import pandas as pd
from openpyxl import load_workbook
from datetime import datetime

# File paths - update these to your actual file paths
registration_file = r"C:\Users\nerme\Desktop\Rnker\Automation\Tournaments Ranking\Austria  23012026 1286_Senior Division.xlsx"
rankings_file = r"C:\Users\nerme\Desktop\Rnker\Automation\Tournaments Ranking\Feb Ranking Seniors.xlsx"
output_file = r"C:\Users\nerme\Desktop\Rnker\Automation\Tournaments Ranking\Austria  23012026 1286_Senior Division With Rankings.xlsx"

# Read rankings file
df_rankings = pd.read_excel(rankings_file)

# Column names - CORRECTED based on actual data
id_col_reg = 'Id'  # Registration file has 'Id' (lowercase 'd')
id_col_rank = 'ID'  # Rankings file has 'ID' (uppercase)
weight_col_rank = 'Weight Category'
# The ranking column is a datetime object, we need to find it
rank_col = [col for col in df_rankings.columns if isinstance(col, datetime)][0]

print(f"Using ranking column: {rank_col}")

# Normalize rankings data
df_rankings['id_normalized'] = df_rankings[id_col_rank].astype(str).str.upper().str.strip()
df_rankings['weight_normalized'] = df_rankings[weight_col_rank].astype(str).str.lower().str.strip()

# Read all sheets from registration file
registration_sheets = pd.read_excel(registration_file, sheet_name=None)

print(f"\nFound {len(registration_sheets)} sheets in registration file:")
print(f"Sheet names: {list(registration_sheets.keys())}\n")

# Load workbook to write back with formulas preserved
wb = load_workbook(registration_file)

total_matched = 0
total_players = 0

# Process each sheet (each weight category)
for sheet_name, df_reg in registration_sheets.items():
    print(f"\nProcessing sheet: {sheet_name}")
    print(f"Players in this category: {len(df_reg)}")
    
    # Get the weight category from sheet name
    weight_category = sheet_name.lower().strip()
    
    # Normalize IDs in registration
    df_reg['id_normalized'] = df_reg[id_col_reg].astype(str).str.upper().str.strip()
    
    # Match with rankings based on ID and weight category
    df_merged = df_reg.merge(
        df_rankings[df_rankings['weight_normalized'] == weight_category][['id_normalized', rank_col]],
        on='id_normalized',
        how='left'
    )
    
    # Count matches
    matched = df_merged[rank_col].notna().sum()
    total_matched += matched
    total_players += len(df_reg)
    
    print(f"Matched: {matched}/{len(df_reg)}")
    
    # Add Ranking column to the Excel sheet
    ws = wb[sheet_name]
    
    # Find the next empty column for Ranking
    max_col = ws.max_column
    ranking_col = max_col + 1
    
    # Add header
    ws.cell(row=1, column=ranking_col, value='Ranking')
    
    # Add ranking values - skip 0 and NaN values
    for idx, rank_value in enumerate(df_merged[rank_col], start=2):
        if pd.notna(rank_value) and rank_value != 0:
            ws.cell(row=idx, column=ranking_col, value=float(rank_value))
        else:
            ws.cell(row=idx, column=ranking_col, value=None)  # Leave blank for 0 or no match

# Save the updated workbook
wb.save(output_file)

print(f"\n{'='*50}")
print(f"✓ Updated file saved as: {output_file}")
print(f"Total players matched: {total_matched}/{total_players}")
print(f"{'='*50}")